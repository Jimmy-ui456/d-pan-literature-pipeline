from __future__ import annotations

import csv
import html
import json
import os
import re
import uuid
from pathlib import Path
from typing import Any, Iterable

from .core import subsection_summary_cn
from .models import AuditEvent, Record


TABLE_HEADERS = [
    "序号",
    "文献层级",
    "标题",
    "期刊",
    "年份及期卷号",
    "通讯作者",
    "方法策略与对应成果",
    "产品",
    "产量核对状态（非自动结论）",
    "证据页码",
    "DOI",
    "筛选理由",
    "可靠性与参考价值",
]


def excel_safe(value: Any) -> Any:
    """Prevent untrusted metadata from becoming a spreadsheet formula."""

    if not isinstance(value, str):
        return value
    value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]", "", value)
    stripped = value.lstrip()
    if stripped.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _atomic_write_text(path: Path, text: str, *, encoding: str = "utf-8") -> None:
    temporary = path.with_name(path.name + f".{uuid.uuid4().hex}.part")
    temporary.write_text(text, encoding=encoding)
    os.replace(temporary, path)


def write_records_json(path: Path, records: list[Record]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    _atomic_write_text(
        path,
        json.dumps([record.to_dict() for record in records], ensure_ascii=False, indent=2) + "\n",
    )


def write_run_audit(
    out_dir: Path,
    events: list[AuditEvent],
    manifest: dict[str, Any],
) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    manifest["event_summary"] = {
        "success": sum(event.status == "success" for event in events),
        "failed": sum(event.status == "failed" for event in events),
        "rejected": sum(event.status == "rejected" for event in events),
        "skipped": sum(event.status == "skipped" for event in events),
    }
    _atomic_write_text(
        out_dir / "run_manifest.json",
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
    )
    _atomic_write_text(
        out_dir / "events.jsonl",
        "".join(json.dumps(event.to_dict(), ensure_ascii=False) + "\n" for event in events),
    )


def _write_csv(path: Path, rows: Iterable[dict[str, Any]], fieldnames: list[str]) -> None:
    temporary = path.with_name(path.name + f".{uuid.uuid4().hex}.part")
    with temporary.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({key: excel_safe(value) for key, value in row.items()})
    os.replace(temporary, path)


def year_volume_issue(record: Record) -> str:
    citation = record.volume
    if record.issue:
        citation += f"({record.issue})"
    if record.pages:
        citation += f":{record.pages}"
    return " ".join(part for part in (record.year, citation) if part)


def _table_rows(records: list[Record]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for index, record in enumerate(records, 1):
        rows.append(
            {
                "序号": index,
                "文献层级": record.library_tier,
                "标题": record.title,
                "期刊": record.journal,
                "年份及期卷号": year_volume_issue(record),
                "通讯作者": record.corresponding_author,
                "方法策略与对应成果": record.method_result_cn,
                "产品": record.product,
                "产量核对状态（非自动结论）": record.final_yield,
                "证据页码": ",".join(map(str, record.results_pages)),
                "DOI": record.doi,
                "筛选理由": "；".join(record.screening_reasons),
                "可靠性与参考价值": (
                    record.reliability_cn
                    + (f" 警告：{'；'.join(record.processing_warnings)}" if record.processing_warnings else "")
                    + f" 来源：{'; '.join(record.source_ids[:6])}"
                ),
            }
        )
    return rows


def _write_workbook(rows: list[dict[str, Any]], output: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "文献总结"
    sheet.append(TABLE_HEADERS)
    for row in rows:
        sheet.append([excel_safe(row.get(header, "")) for header in TABLE_HEADERS])
    widths = [8, 19, 48, 26, 18, 34, 72, 20, 32, 14, 30, 52, 58]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    temporary = output.with_name(output.stem + f".{uuid.uuid4().hex}.part" + output.suffix)
    workbook.save(temporary)
    os.replace(temporary, output)


def result_subsections(text: str, max_sections: int = 12) -> list[tuple[str, str]]:
    if not text:
        return []
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    sections: list[tuple[str, list[str]]] = []
    current_title = "Results / Results and Discussion"
    current_body: list[str] = []
    heading_re = re.compile(
        r"^((?:\d+(?:\.\d+){0,3}|[A-Z])\.?\s+)?"
        r"([A-Z][A-Za-z0-9βαγδµμ/(),:;+\-\s]{8,140})$"
    )
    for line in lines:
        if line.startswith("[[PAGE "):
            current_body.append(line)
            continue
        is_heading = bool(heading_re.match(line)) and len(line.split()) <= 18 and not line.endswith(".")
        if is_heading and current_body:
            sections.append((current_title, current_body))
            current_title = line
            current_body = []
            if len(sections) >= max_sections:
                break
        else:
            current_body.append(line)
    if current_body and len(sections) < max_sections:
        sections.append((current_title, current_body))
    return [(title, " ".join(body)[:1_800]) for title, body in sections if body]


def _md_text(value: str) -> str:
    escaped = html.escape(value, quote=False)
    escaped = escaped.replace("\\", "\\\\").replace("\r", " ").replace("\n", " ")
    return re.sub(r"([`*_{}\[\]()#!|~])", r"\\\1", escaped)


def _digest(records: list[Record]) -> str:
    lines = ["# 课题文献候选证据定位（不是综述结论）", ""]
    lines.append("> 自动提取只用于定位，不能替代人工阅读；页码以 PDF 页序为准。")
    for index, record in enumerate(records, 1):
        lines.extend(
            [
                "",
                f"## {index}. {_md_text(record.title)}",
                "",
                f"- DOI: {_md_text(record.doi or '无')}",
                f"- PDF: {_md_text(record.pdf_file or '未下载')}",
                f"- 证据页码: {','.join(map(str, record.results_pages)) or '未定位'}",
                f"- 筛选理由: {_md_text('；'.join(record.screening_reasons))}",
            ]
        )
        if record.results_excerpt:
            for section_index, (heading, body) in enumerate(result_subsections(record.results_excerpt), 1):
                safe_body = body.replace("```", "'''")
                lines.extend(
                    [
                        "",
                        f"### {index}.{section_index} {_md_text(heading)}",
                        "",
                        f"- 中文提炼：{_md_text(subsection_summary_cn(record, heading, body))}",
                        "",
                        "```text",
                        safe_body[:1_200],
                        "```",
                    ]
                )
        else:
            lines.extend(
                [
                    "",
                    "### 未可靠定位结果段",
                    "",
                    f"- 中文提炼：{_md_text(record.method_result_cn)}",
                ]
            )
    return "\n".join(lines) + "\n"


def write_screening_log(records: list[Record], retained: list[Record], out_dir: Path) -> None:
    retained_ids = {id(record) for record in retained}
    rows = []
    for record in records:
        rows.append(
            {
                "标题": record.title,
                "DOI": record.doi,
                "分数": record.relevance_score,
                "文献层级": record.library_tier,
                "是否纳入": "是" if id(record) in retained_ids else "否",
                "筛选理由": "；".join(record.screening_reasons),
            }
        )
    _write_csv(
        out_dir / "screening_decisions.csv",
        rows,
        ["标题", "DOI", "分数", "文献层级", "是否纳入", "筛选理由"],
    )


def write_outputs(
    records: list[Record],
    out_dir: Path,
    *,
    all_screened_records: list[Record] | None = None,
    manifest: dict[str, Any] | None = None,
    events: list[AuditEvent] | None = None,
) -> list[AuditEvent]:
    out_dir.mkdir(parents=True, exist_ok=True)
    output_events = list(events or [])
    write_records_json(out_dir / "literature_records.json", records)
    rows = _table_rows(records)
    _write_csv(
        out_dir / "literature_candidates.csv",
        rows,
        TABLE_HEADERS,
    )
    coverage_rows = [
        {
            "序号": index,
            "标题": record.title,
            "DOI": record.doi,
            "PDF状态": "已下载/已关联" if record.pdf_file else "未取得可验证PDF",
            "PDF路径": record.pdf_file,
            "Results解析状态": "已定位结果段" if record.results_excerpt else "未可靠定位结果段",
            "证据页码": ",".join(map(str, record.results_pages)),
            "备注": record.reliability_cn,
        }
        for index, record in enumerate(records, 1)
    ]
    _write_csv(
        out_dir / "download_coverage.csv",
        coverage_rows,
        ["序号", "标题", "DOI", "PDF状态", "PDF路径", "Results解析状态", "证据页码", "备注"],
    )
    try:
        _write_workbook(rows, out_dir / "literature_candidates.xlsx")
    except (ImportError, OSError, ValueError) as exc:
        output_events.append(
            AuditEvent(
                stage="output",
                status="failed",
                source="openpyxl",
                message=f"{type(exc).__name__}: {exc}",
            )
        )
    _atomic_write_text(out_dir / "results_digest.md", _digest(records))
    if all_screened_records is not None:
        write_records_json(out_dir / "all_screened_records.json", all_screened_records)
        write_screening_log(all_screened_records, records, out_dir)
    if manifest is not None:
        write_run_audit(out_dir, output_events, manifest)
    else:
        _atomic_write_text(
            out_dir / "events.jsonl",
            "".join(json.dumps(event.to_dict(), ensure_ascii=False) + "\n" for event in output_events),
        )
    return output_events
